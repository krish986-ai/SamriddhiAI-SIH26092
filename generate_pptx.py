import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Colors - Warm Indian Saffron / Orange Theme
    c_navy = RGBColor(17, 24, 39)        # #111827
    c_dark_bg = RGBColor(15, 23, 42)     # #0f172a
    c_card_bg = RGBColor(30, 41, 59)     # #1e293b
    c_orange = RGBColor(234, 88, 12)     # #ea580c (BHUSEWA Orange)
    c_bright_orange = RGBColor(249, 115, 22) # #f97316
    c_amber = RGBColor(245, 158, 11)     # #f59e0b
    c_purple = RGBColor(88, 80, 236)     # #5850ec (ENGORIO purple)
    c_white = RGBColor(255, 255, 255)
    c_light_gray = RGBColor(203, 213, 225)

    def add_header(slide, title, category="SMART INDIA HACKATHON 2026 | PROBLEM SIH26092"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_bright_orange
        
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = c_white

    def create_card(slide, left, top, width, height, title, content_list, border_color=c_orange, bg_color=c_card_bg):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_head = tf.paragraphs[0]
        p_head.text = title
        p_head.font.size = Pt(14)
        p_head.font.bold = True
        p_head.font.color.rgb = border_color
        
        for item in content_list:
            p = tf.add_paragraph()
            p.text = f"• {item}"
            p.font.size = Pt(11)
            p.font.color.rgb = c_light_gray
            p.space_before = Pt(6)

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_dark_bg
    bg1.line.fill.background()

    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "SMART INDIA HACKATHON 2026 | PROBLEM STATEMENT SIH26092"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = c_bright_orange
    p.space_after = Pt(10)

    p = tf1.add_paragraph()
    p.text = "BHUSEWA AI"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = c_white

    p = tf1.add_paragraph()
    p.text = "Next-Gen AI-Driven Concessional Scheme Matching, Financial Triage & SCA Dispatch Platform"
    p.font.size = Pt(18)
    p.font.color.rgb = c_bright_orange
    p.space_after = Pt(25)

    p = tf1.add_paragraph()
    p.text = "Aligned with Ministry of Social Justice and Empowerment (MoSJE) & myScheme.gov.in"
    p.font.size = Pt(13)
    p.font.color.rgb = c_light_gray

    p = tf1.add_paragraph()
    p.text = "Developed by Team Innovision | Engineered by ENGORIO"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = c_purple
    p.space_before = Pt(15)

    # Slide 2: Problem Statement & Vision
    s2 = prs.slides.add_slide(blank_slide_layout)
    bg2 = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = c_dark_bg
    bg2.line.fill.background()
    add_header(s2, "Problem Statement & Solution Vision")

    create_card(s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8), 
                "1. Critical Challenges", 
                [
                    "Information Fragmentation: Citizens struggle across 4,770+ central & state schemes.",
                    "Low Concessional Awareness: Beneficiaries take 24%+ market loans unaware of 4-6% apex rates.",
                    "Language & Literacy Barriers: Complex legal text hinders rural & vernacular citizens.",
                    "SCA Application Bottlenecks: Manual document triage causes 40%+ delay and rejections."
                ], border_color=c_amber)

    create_card(s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8), 
                "2. The BHUSEWA AI Solution", 
                [
                    "Instant Multi-Criteria AI Matching: 100% accurate rule scoring across SC/ST/OBC criteria.",
                    "Live Govt Ingestion Pipeline: Auto-syncs new central schemes from data.gov.in & myScheme.",
                    "BhuSewa AI Assistant: Floating multilingual voice & chat assistant in 4 languages.",
                    "Pre-Vetted SCA Routing: Direct geo-spatial dispatch to 36 State Channelizing Agencies."
                ], border_color=c_orange)

    create_card(s2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8), 
                "3. Key Impact Goals", 
                [
                    "Zero Commission Fees: 100% free direct-to-government welfare dispatch.",
                    "Sub-300ms Evaluation: Real-time calculation of interest subsidies (up to 35% DBT).",
                    "GIGW 3.0 Compliance: Full 12-tile accessibility suite by ENGORIO.",
                    "End-to-End Tracking: SMS and transparent stage-by-stage status tracker."
                ], border_color=c_purple)

    # Slide 3: System Architecture
    s3 = prs.slides.add_slide(blank_slide_layout)
    bg3 = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg3.fill.solid()
    bg3.fill.fore_color.rgb = c_dark_bg
    bg3.line.fill.background()
    add_header(s3, "System Architecture & End-to-End Data Pipeline")

    create_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
                "Frontend & Design System",
                [
                    "React 18 + Vite: High performance sub-500ms reactive client rendering.",
                    "BHUSEWA Orange Palette: Official myScheme layout with saffron/orange theme.",
                    "ENGORIO GIGW 3.0 Accessibility Suite: 12-tile assistive toolset engineered by ENGORIO."
                ], border_color=c_purple)

    create_card(s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3),
                "AI Inference & Triage Engine",
                [
                    "Deterministic Rule Scoring Engine: NSFDC, NBCFDC, Stand-Up India, PMEGP, MoSJE logic.",
                    "Multi-Variable Scoring: Evaluates social caste, income limit, gender, project cost, location.",
                    "Zero Hallucination: Grounded mathematical calculations for concessional EMIs."
                ], border_color=c_orange)

    create_card(s3, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.3),
                "Automated Govt Sync Pipeline",
                [
                    "GovernmentSchemeSyncEngine: Automated REST fetcher for data.gov.in & myScheme API.",
                    "Intelligent Schema Normalizer: Cleans, deduplicates, and structures live scheme parameters.",
                    "Hot-Reload Integration: Synced schemes instantly participate in live AI citizen triage."
                ], border_color=c_amber)

    create_card(s3, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.3),
                "Voice & Multimodal Interface",
                [
                    "Web Speech STT & SpeechSynthesis: Voice guidance in English, Hindi, Marathi, Tamil.",
                    "BhuSewa AI Agent: Context-aware conversational assistant with deep-link navigation.",
                    "Direct SCA Dispatch Portal: Officer verification workflow and triage queues."
                ], border_color=c_orange)

    # Slide 4: Automated Govt Scheme Fetcher
    s4 = prs.slides.add_slide(blank_slide_layout)
    bg4 = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg4.fill.solid()
    bg4.fill.fore_color.rgb = c_dark_bg
    bg4.line.fill.background()
    add_header(s4, "Automated Government Scheme Fetcher & Ingestion Engine")

    create_card(s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
                "1. Data Acquisition",
                [
                    "Official Endpoints: Connects to data.gov.in & api.myscheme.gov.in.",
                    "Scheduled Cron & On-Demand: Daily automated polling + manual officer trigger.",
                    "Resilient Fallbacks: Handles network interruptions with persistent local cache."
                ], border_color=c_orange)

    create_card(s4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
                "2. Schema Normalizer",
                [
                    "Standardization: Normalizes raw government fields into standardized JSON format.",
                    "Field Extraction: Extracts nodal agency, interest rates, capital subsidies, caste criteria.",
                    "Deduplication: Compares scheme hashes to prevent redundant records."
                ], border_color=c_amber)

    create_card(s4, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
                "3. Real-Time Deployment",
                [
                    "Active State Injection: Ingested schemes immediately appear in 4,770+ database.",
                    "Instant AI Scoring: Citizen profile matching updates with zero code deployments.",
                    "Live Ingested Examples: PM-AJAY 4.5% grants & NBCFDC Green Business 4% loans."
                ], border_color=c_purple)

    # Slide 5: ENGORIO Accessibility Suite
    s5 = prs.slides.add_slide(blank_slide_layout)
    bg5 = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg5.fill.solid()
    bg5.fill.fore_color.rgb = c_dark_bg
    bg5.line.fill.background()
    add_header(s5, "ENGORIO Accessibility Suite (GIGW 3.0 Compliant)")

    create_card(s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
                "12 Assistive Tools (3x4 Grid)",
                [
                    "1. Bigger Text / 2. Smaller Text: Responsive scaling from 80% to 130%.",
                    "3. Text Spacing: Letter & word spacing adjustments for reading clarity.",
                    "4. Line Height: 4 progressive line spacing modes with visual check indicators.",
                    "5. Dyslexia Friendly: OpenDyslexic high-legibility font formatting.",
                    "6. ADHD Mode: Cursor-tracking reading spotlight ruler.",
                    "7. Saturation / 8. Invert Colors: High/low saturation and color inversion.",
                    "9. Highlight Links / 10. Large Cursor / 11. Pause Animations / 12. Hide Images."
                ], border_color=c_purple)

    create_card(s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
                "Key Usability Innovations",
                [
                    "Right-Edge Floating Trigger: Minimalist universal accessibility icon button on right edge.",
                    "Global Keyboard Shortcut: Press Ctrl+F2 anywhere on the platform to launch options.",
                    "State Persistence & 1-Click Reset: Instant reset button restoring default view.",
                    "Optimized Dark Theme: Rich slate palette (#0b1120) with orange accents & glowing tags.",
                    "Official Branding: Authenticated with 'Created by ENGORIO' badge."
                ], border_color=c_orange)

    # Slide 6: Citizen Tools
    s6 = prs.slides.add_slide(blank_slide_layout)
    bg6 = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg6.fill.solid()
    bg6.fill.fore_color.rgb = c_dark_bg
    bg6.line.fill.background()
    add_header(s6, "Citizen Tools: Concessional Calculator & SCA Locator")

    create_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
                "Concessional Financial Calculator",
                [
                    "Apex Comparison: Compares 4% NSFDC vs 14% commercial bank EMI rates.",
                    "DBT Capital Subsidy: Calculates up to 35% upfront subsidy under PMEGP.",
                    "Interactive Sliders: Real-time project cost and tenure adjustment.",
                    "Total Savings Highlight: Displays exact Lakhs saved over repayment tenure."
                ], border_color=c_orange)

    create_card(s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
                "Geo-Spatial SCA Partner Locator",
                [
                    "36 States & UTs Coverage: Direct directory of State Channelizing Agencies.",
                    "Interactive Map: Visual clickable state pins with district nodal officers.",
                    "Instant Routing: Dispatches application payload directly to local SCA desk.",
                    "Application Tracker: Real-time stage monitoring from submission to sanction."
                ], border_color=c_amber)

    # Slide 7: Conclusion & Roadmap
    s7 = prs.slides.add_slide(blank_slide_layout)
    bg7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg7.fill.solid()
    bg7.fill.fore_color.rgb = c_dark_bg
    bg7.line.fill.background()
    add_header(s7, "Conclusion, Deliverables & Future Roadmap")

    create_card(s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
                "1. Verified Deliverables",
                [
                    "100% myScheme.gov.in Look & Feel with BHUSEWA Orange theme.",
                    "Sub-300ms AI Triage Engine.",
                    "Live Government Scheme Sync Engine.",
                    "Multilingual BhuSewa AI Assistant.",
                    "ENGORIO GIGW 3.0 Suite by ENGORIO."
                ], border_color=c_orange)

    create_card(s7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
                "2. National Impact",
                [
                    "Promotes affirmative credit access.",
                    "Prevents loan-shark exploitation.",
                    "Reduces SCA processing from 4 weeks to 3 days.",
                    "Zero commission for marginalized citizens."
                ], border_color=c_amber)

    create_card(s7, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
                "3. Future Roadmap",
                [
                    "DigiLocker & Aadhaar e-KYC integration.",
                    "WhatsApp & Telegram bot for rural access.",
                    "Automated SMS alerts on scheme sanction.",
                    "Direct DBT gateway linkage."
                ], border_color=c_purple)

    # Save to Desktop
    desktop_path = os.path.join(os.path.expanduser("~"), "Desktop")
    pptx_output = os.path.join(desktop_path, "BHUSEWA_Project_Presentation.pptx")
    prs.save(pptx_output)
    print(f"Presentation saved successfully to {pptx_output}")

if __name__ == "__main__":
    create_presentation()
